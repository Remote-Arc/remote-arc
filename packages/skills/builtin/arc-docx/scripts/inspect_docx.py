"""Read-only DOCX structure and content inspection; not a page renderer."""
import argparse
import json
from pathlib import Path
from docx import Document


def inspect(filename):
    document = Document(filename)
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    return {
        "path": str(Path(filename).resolve()),
        "paragraphs": len(paragraphs),
        "headings": [{"style": p.style.name, "text": p.text} for p in paragraphs if p.style.name.startswith("Heading")],
        "text": "\n".join(p.text for p in paragraphs),
        "tables": [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables],
        "sections": [{"width_in": round(s.page_width.inches, 3), "height_in": round(s.page_height.inches, 3)} for s in document.sections],
        "verification": "structure_only_not_visual_layout",
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    args = parser.parse_args()
    print(json.dumps(inspect(args.file), ensure_ascii=False, indent=2))
