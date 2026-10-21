"""Inspect PPTX text and top-level canvas overflow; does not render slides."""
import argparse
import json
from pathlib import Path
from pptx import Presentation


def inspect(filename):
    deck = Presentation(filename)
    slides = []
    for index, slide in enumerate(deck.slides, 1):
        shapes = []
        for shape in slide.shapes:
            outside = shape.left < 0 or shape.top < 0 or shape.left + shape.width > deck.slide_width or shape.top + shape.height > deck.slide_height
            shapes.append({"name": shape.name, "text": shape.text if shape.has_text_frame else "", "outside_canvas": outside})
        slides.append({"number": index, "shapes": shapes})
    return {"path": str(Path(filename).resolve()), "slide_count": len(slides), "slides": slides, "verification": "top_level_geometry_only_not_visual_layout"}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    print(json.dumps(inspect(parser.parse_args().file), ensure_ascii=False, indent=2))
