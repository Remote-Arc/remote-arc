"""Inspect XLSX formulas, stored error cells, and cached values without changing it."""
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook


def inspect(filename):
    formulas = load_workbook(filename, read_only=True, data_only=False)
    cached = load_workbook(filename, read_only=True, data_only=True)
    try:
        sheets = []
        for sheet in formulas:
            result = {"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column, "formulas": [], "stored_errors": []}
            for row, cached_row in zip(sheet.iter_rows(), cached[sheet.title].iter_rows()):
                for cell, value in zip(row, cached_row):
                    if cell.data_type == "f":
                        result["formulas"].append({"cell": cell.coordinate, "formula": cell.value, "cached": value.value, "cache_missing": value.value is None})
                    if cell.data_type == "e" or value.data_type == "e":
                        result["stored_errors"].append({"cell": cell.coordinate, "error": value.value if value.data_type == "e" else cell.value})
            sheets.append(result)
        return {"path": str(Path(filename).resolve()), "sheets": sheets, "recalculated": False}
    finally:
        formulas.close()
        cached.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    print(json.dumps(inspect(parser.parse_args().file), ensure_ascii=False, indent=2))
