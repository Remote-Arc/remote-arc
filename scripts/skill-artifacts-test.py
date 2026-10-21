"""Exercise actual office artifacts and data runtimes without accounts or paid APIs."""
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "work" / "skill-artifacts-qa"
OUT.mkdir(parents=True, exist_ok=True)


def inspect(skill, helper, filename):
    before = hashlib.sha256(filename.read_bytes()).hexdigest()
    result = subprocess.run([sys.executable, str(ROOT / "packages" / "skills" / "builtin" / skill / "scripts" / helper), str(filename)], check=True, capture_output=True, text=True, encoding="utf-8", env={**os.environ, "PYTHONUTF8": "1", "PYTHONDONTWRITEBYTECODE": "1"})
    assert before == hashlib.sha256(filename.read_bytes()).hexdigest(), "Inspector modified the artifact"
    return json.loads(result.stdout)


from docx import Document
document = Document()
document.add_heading("Quarterly evidence", 1)
document.add_paragraph("Revenue verified from supplied data.")
document.add_table(rows=2, cols=2).cell(1, 1).text = "1200"
docx = OUT / "report.docx"
document.save(docx)
word = inspect("arc-docx", "inspect_docx.py", docx)
assert word["headings"][0]["text"] == "Quarterly evidence"
assert word["tables"][0][1][1] == "1200"
assert word["verification"] == "structure_only_not_visual_layout"

from openpyxl import Workbook
workbook = Workbook()
sheet = workbook.active
sheet.title = "Summary"
sheet.append(["Revenue", "Cost", "Margin"])
sheet.append([1200, 300, "=A2-B2"])
sheet["D2"] = "#REF!"
xlsx = OUT / "budget.xlsx"
workbook.save(xlsx)
excel = inspect("arc-xlsx", "inspect_xlsx.py", xlsx)
assert excel["sheets"][0]["formulas"][0]["cache_missing"] is True
assert excel["sheets"][0]["stored_errors"][0]["error"] == "#REF!"
assert excel["recalculated"] is False

from pptx import Presentation
from pptx.util import Inches
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[6])
slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Verified release"
slide.shapes.add_textbox(-Inches(1), 0, Inches(1), Inches(1)).text = "Intentional overflow fixture"
pptx = OUT / "deck.pptx"
deck.save(pptx)
powerpoint = inspect("arc-pptx", "inspect_pptx.py", pptx)
assert powerpoint["slide_count"] == 1
assert any(shape["outside_canvas"] for shape in powerpoint["slides"][0]["shapes"])

from pypdf import PdfWriter
writer = PdfWriter()
writer.add_blank_page(width=200, height=300)
writer.add_blank_page(width=400, height=500).rotate(90)
pdf = OUT / "ordered.pdf"
writer.write(pdf)
pages = inspect("arc-pdf", "inspect_pdf.py", pdf)
assert pages["page_count"] == 2
assert pages["pages"][1]["width_pt"] == 400
assert pages["pages"][1]["rotation"] == 90
assert pages["pages"][0]["has_extractable_text"] is False
writer.encrypt("fixture-password")
locked = OUT / "encrypted.pdf"
writer.write(locked)
assert inspect("arc-pdf", "inspect_pdf.py", locked)["locked"] is True

import matplotlib
matplotlib.use("Agg")
import seaborn as sns
import matplotlib.pyplot as plt
import pandas as pd
data = pd.DataFrame({"city": ["Auckland", "Wellington", "Auckland", "Wellington"], "revenue": [739, 842, 811, 905]})
sns.barplot(data=data, x="city", y="revenue", errorbar=None)
chart = OUT / "revenue.png"
plt.savefig(chart, dpi=120, bbox_inches="tight")
plt.close()
assert chart.read_bytes().startswith(b"\x89PNG\r\n\x1a\n")

import networkx as nx
graph = nx.Graph([("brief", "draft"), ("draft", "review"), ("review", "publish")])
assert nx.shortest_path(graph, "brief", "publish") == ["brief", "draft", "review", "publish"]
assert nx.number_connected_components(graph) == 1

from sklearn.datasets import load_iris
from sklearn.model_selection import train_test_split
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.linear_model import LogisticRegression
x, y = load_iris(return_X_y=True)
x_train, x_test, y_train, y_test = train_test_split(x, y, test_size=0.25, random_state=19, stratify=y)
model = make_pipeline(StandardScaler(), LogisticRegression(max_iter=300))
model.fit(x_train, y_train)
assert model.score(x_test, y_test) > 0.85
assert model[0].n_samples_seen_ == len(x_train), "Preprocessing leaked test data"

import numpy as np
helper = ROOT / "packages/skills/vendor/community-statistical-analysis/scripts/assumption_checks.py"
spec = importlib.util.spec_from_file_location("skill_assumption_checks", helper)
diagnostics = importlib.util.module_from_spec(spec)
spec.loader.exec_module(diagnostics)
normality = diagnostics.check_normality(np.random.default_rng(19).normal(size=120), plot=False)
assert np.isfinite(normality["p_value"])
try:
    diagnostics.check_normality([1, 1, 1], plot=False)
except ValueError:
    pass
else:
    raise AssertionError("Constant data must not manufacture a valid normality test")

print(json.dumps({"passed": True, "office_artifacts": 4, "read_only_inspectors": True, "missing_formula_cache_detected": True, "stored_formula_errors_detected": True, "slide_overflow_detected": True, "encrypted_pdf_reported": True, "seaborn_png": str(chart), "networkx": True, "sklearn_holdout_score": model.score(x_test, y_test), "actual_statistical_helper": True}, indent=2))
